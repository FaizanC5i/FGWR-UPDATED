from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from io import BytesIO
import os

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

class ChartData(BaseModel):
    labels: List[str]
    values: List[float]

class TrendChartData(BaseModel):
    labels: List[str]
    damages: List[float]
    stales: List[float]

class ReportPayload(BaseModel):
    s1_request: str
    s1_generated_on: str 
    s2_body: List[str]
    s2_kpi1_label: str 
    s2_kpi2_label: str
    s2_kpi3_label: str
    s2_kpi1_value: str
    s2_kpi2_value: str
    s2_kpi3_value: str
    s3_insights: List[str]
    s3_chart_trend: TrendChartData
    s4_body_bullets: List[str] 
    s4_chart_top_drivers: ChartData
    s5_recommendations: List[str]

def replace_text_keep_format(shape, new_text):
    """Replaces text in standard boxes like KPIs or Headers, keeping format."""
    if not shape.has_text_frame:
        return
    text_frame = shape.text_frame
    if text_frame.paragraphs and text_frame.paragraphs[0].runs:
        text_frame.paragraphs[0].runs[0].text = str(new_text)
        for run in text_frame.paragraphs[0].runs[1:]:
            run.text = ""
    else:
        shape.text = str(new_text)

def replace_bullets_keep_format(shape, lines: List[str]):
    """
    Safely injects an array of text into native PPT bullets.
    It clones the font settings (color, size, weight) of the first master bullet.
    """
    if not shape.has_text_frame or not lines:
        return
    
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    
    p0 = tf.paragraphs[0]
    
    # Extract formatting from the template's first placeholder bullet
    font_name, font_size, font_bold, font_italic, font_color_rgb, font_color_theme = (None,)*6
    if p0.runs:
        f = p0.runs[0].font
        font_name = f.name
        font_size = f.size
        font_bold = f.bold
        font_italic = f.italic
        if f.color and f.color.type:
            try: font_color_rgb = f.color.rgb
            except: pass
            try: font_color_theme = f.color.theme_color
            except: pass

    # Clear template dummy text, keeping the box-level bullet configuration
    tf.clear() 

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0 # Ensures it is a top-level bullet
        
        # Re-apply the extracted font formatting to the new line
        if p.runs:
            f = p.runs[0].font
            if font_name: f.name = font_name
            if font_size: f.size = font_size
            if font_bold is not None: f.bold = font_bold
            if font_italic is not None: f.italic = font_italic
            if font_color_rgb: 
                f.color.rgb = font_color_rgb
            elif font_color_theme: 
                f.color.theme_color = font_color_theme

@app.post("/api/generate-ppt")
async def generate_ppt(payload: ReportPayload):
    current_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(current_dir, "C5i_Auto_Report_Template_v1.pptx")
    prs = Presentation(template_path)
    
    for slide in prs.slides:
        for shape in slide.shapes:
            
            # Use .lower() to completely avoid case-sensitivity bugs with templates
            s_name = shape.name.lower()
            
            # --- SLIDE 1 ---
            if s_name == "s1_request":
                replace_text_keep_format(shape, f"Report Request: {payload.s1_request}")
            elif s_name == "s1_generatedon":
                replace_text_keep_format(shape, payload.s1_generated_on)
                
            # --- SLIDE 2 ---
            elif s_name == "s2_kpi1_label":
                replace_text_keep_format(shape, payload.s2_kpi1_label)
            elif s_name == "s2_kpi2_label":
                replace_text_keep_format(shape, payload.s2_kpi2_label)
            elif s_name == "s2_kpi3_label":
                replace_text_keep_format(shape, payload.s2_kpi3_label)
            elif s_name == "s2_kpi1_value":
                replace_text_keep_format(shape, payload.s2_kpi1_value)
            elif s_name == "s2_kpi2_value":
                replace_text_keep_format(shape, payload.s2_kpi2_value)
            elif s_name == "s2_kpi3_value":
                replace_text_keep_format(shape, payload.s2_kpi3_value)
            elif s_name == "s2_body":
                replace_bullets_keep_format(shape, payload.s2_body)

            # --- SLIDE 3 ---
            elif s_name == "s3_insights":
                replace_bullets_keep_format(shape, payload.s3_insights)
            elif s_name == "s3_chart_trend" and shape.has_chart:
                chart_data = CategoryChartData()
                chart_data.categories = payload.s3_chart_trend.labels
                chart_data.add_series('Damages', tuple(payload.s3_chart_trend.damages))
                chart_data.add_series('Stales', tuple(payload.s3_chart_trend.stales))
                shape.chart.replace_data(chart_data)

            # --- SLIDE 4 ---
            elif s_name == "s4_body_bullets":
                replace_bullets_keep_format(shape, payload.s4_body_bullets)
            elif s_name == "s4_chart_top_drivers" and shape.has_chart:
                chart_data = CategoryChartData()
                chart_data.categories = payload.s4_chart_top_drivers.labels
                chart_data.add_series('Drivers', tuple(payload.s4_chart_top_drivers.values))
                shape.chart.replace_data(chart_data)

            # --- SLIDE 5 ---
            # Catches both naming conventions you might have used
            elif s_name in ["s5_recommendations", "s5_body"]: 
                replace_bullets_keep_format(shape, payload.s5_recommendations)

    output = BytesIO()
    prs.save(output)
    output.seek(0)

    return StreamingResponse(
        output, 
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": "attachment; filename=C5i_Executive_Report.pptx"}
    )